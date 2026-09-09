import re
import zipfile
from collections import Counter

import openpyxl

import config
import pipeline
import write_output


def _spelling_occurrence_counts(original_rows):
    """Re-derive, from the ORIGINAL (pre-cleaning) text, exactly how many
    times each known misspelling occurred - used for the Change Log sheet
    and the console summary."""
    all_text = " ".join(str(r["Defects Description"] or "") for r in original_rows)
    counts = []
    for bad, good in config.SPELLING_MAP.items():
        n = len(re.findall(r"\b" + re.escape(bad) + r"\b", all_text, re.I))
        if n > 0:
            counts.append((bad, good, n))
    return counts


def _categorical_change_counts(original_rows):
    out = []
    for col, mapping in config.CATEGORICAL_MAP.items():
        c = Counter()
        for r in original_rows:
            v = r[col]
            if v is None:
                continue
            key = str(v).strip().lower()
            if key in mapping and mapping[key] != v:
                c[(v, mapping[key])] += 1
        for (before, after), cnt in c.items():
            out.append((col, before, after, cnt))
    return out


def main():
    print(f"Loading source workbook: {config.SOURCE_FILE}")
    try:
        readable_copy = pipeline.get_readable_copy(config.SOURCE_FILE)
        wb = openpyxl.load_workbook(readable_copy, data_only=True)
        if config.MASTER_SHEET not in wb.sheetnames:
            raise pipeline.PipelineInputError(
                f"The required Master sheet '{config.MASTER_SHEET}' was not found."
            )
        month_sheets = config.detect_month_sheets(wb.sheetnames)
        if not month_sheets:
            raise pipeline.PipelineInputError(
                "No monthly sheets were found. Expected names such as 'FEBRUARY 2026'."
            )
    except FileNotFoundError:
        print(f"ERROR: Source workbook was not found: {config.SOURCE_FILE}")
        return 1
    except (PermissionError, OSError, ValueError, zipfile.BadZipFile,
            openpyxl.utils.exceptions.InvalidFileException) as exc:
        print(f"ERROR: Could not open the source workbook '{config.SOURCE_FILE}': {exc}")
        return 1
    except pipeline.PipelineInputError as exc:
        print(f"ERROR: {exc}")
        return 1

    # --- 1. Read ---
    try:
        master_rows = pipeline.read_master(wb)
    except pipeline.PipelineInputError as exc:
        print(f"ERROR: {exc}")
        return 1
    print(f"  Master rows read: {len(master_rows)}")

    monthly_rows_by_month = {}
    for month, sheet_name in month_sheets:
        try:
            rows = pipeline.read_month(wb, month, sheet_name)
        except pipeline.PipelineInputError as exc:
            print(f"ERROR: {exc}")
            return 1
        monthly_rows_by_month[month] = rows
        print(f"  {month} ({sheet_name}) rows read: {len(rows)}")

    # --- 2. Reconcile (find what's missing, append it) ---
    final_rows, stats = pipeline.reconcile(master_rows, monthly_rows_by_month)
    print(f"\nRows appended: {stats['appended_total']} "
          f"({dict(stats['appended_by_month'])})")
    print(f"Final merged row count: {len(final_rows)}")
    if stats["anomalies_master_has_more"]:
        print(f"  WARNING: {len(stats['anomalies_master_has_more'])} rows where Master has "
              f"MORE copies of a key than the source - review these manually.")
        
    maint_stats = pipeline.fix_maintenance_by_contamination(final_rows)
    print(f"\nMaintenance By contamination: {len(maint_stats['auto_fixed'])} auto-corrected, "
      f"{len(maint_stats['flagged_for_review'])} left blank + flagged (no reliable reference).")

    # keep a copy of the pre-cleaning text for the change-log counts below
    original_rows_snapshot = [dict(r) for r in final_rows]

    # --- 3. Standardize text ---
    motorized_before = sum(
        row["Motorized/Non-Motorized"] == "Motorized" for row in final_rows
    )
    print(f"Motorized rows before standardization: {motorized_before}")
    text_stats = pipeline.standardize_text(final_rows)
    motorized_after = sum(
        row["Motorized/Non-Motorized"] == "Motorized" for row in final_rows
    )
    print(f"Motorized rows after standardization: {motorized_after}")
    if motorized_after != motorized_before:
        raise RuntimeError(
            "Motorized row count changed during standardization: "
            f"{motorized_before} before, {motorized_after} after."
        )
    print(f"\nDefects Description rows changed: {text_stats['desc_rows_fixed']}")
    print(f"Categorical field cells standardized: {text_stats['categorical_fixed']}")

    # DISABLED EQUIPMENT-CODE QA: uncomment with the marked blocks in
    # pipeline.py, write_output.py, and config.py to restore this check.
    # equipment_code_qa = pipeline.check_equipment_code_matches(final_rows)
    # print("Equipment-code QA: "
    #       f"matched={equipment_code_qa['matched']}, "
    #       f"no_code_present={equipment_code_qa['no_code_present']}, "
    #       f"possible_mismatch={equipment_code_qa['possible_mismatch']}")

    # --- 4. Flag Malay text (not translated) ---
    malay_count = pipeline.flag_malay(final_rows)
    print(f"Rows flagged for Malay content (not translated): {malay_count}")

    # --- 5. Convert Date to real date type ---
    date_failures = pipeline.convert_dates(final_rows)
    print(f"Dates converted to true date type. Parse failures: {date_failures}")

    # --- 6. Apply identifier cleaning, audit business-key repeats, and remove
    # only rows that are exact duplicates across the complete canonical row. ---
    equipment_numbers_changed = pipeline.standardize_equipment_numbers(final_rows)
    duplicate_audit = pipeline.duplicate_audit(final_rows)
    final_rows, exact_duplicates_removed = pipeline.remove_exact_duplicates(final_rows)
    stats.update({
        "rows_imported": stats["master_total"] + sum(stats["monthly_totals"].values()),
        "rows_exported": len(final_rows),
        "exact_duplicates_removed": exact_duplicates_removed,
        "duplicate_groups_detected": len(duplicate_audit),
        "duplicate_audit": duplicate_audit,
        "blank_counts": pipeline.blank_counts(final_rows),
        "equipment_numbers_changed": equipment_numbers_changed,
    })
    print(f"Equipment numbers standardized: {equipment_numbers_changed}")
    print(f"Exact duplicate rows removed: {exact_duplicates_removed}")
    print(f"Duplicate groups detected: {len(duplicate_audit)}")

    # --- 7. Build change-log detail for the output workbook ---
    spelling_counts = _spelling_occurrence_counts(original_rows_snapshot)
    categorical_rows = _categorical_change_counts(original_rows_snapshot)
    format_fix_counts = [
        ("\"non negligence\" -> \"non-negligence\" (spacing)",
         sum(1 for r in original_rows_snapshot
             if r["Defects Description"] and re.search(r"\bnon\s+negligence\b", r["Defects Description"], re.I))),
        ("Trailing period(s) removed",
         sum(1 for r in original_rows_snapshot
             if r["Defects Description"] and str(r["Defects Description"]).strip().endswith("."))),
        ("Parenthesis spacing tightened",
         sum(1 for r in original_rows_snapshot
             if r["Defects Description"] and re.search(r"\(\s|\s\)", str(r["Defects Description"])))),
    ]

    # --- 8. Write output workbook ---
    write_output.write_workbook(
        final_rows, stats, text_stats, malay_count, date_failures,
        spelling_counts, categorical_rows, format_fix_counts,
        maint_stats,
        # equipment_code_qa,  # DISABLED EQUIPMENT-CODE QA
        [month for month, _ in month_sheets],
        config.OUTPUT_FILE,
    )
    print(f"\nSaved: {config.OUTPUT_FILE}")


if __name__ == "__main__":
    raise SystemExit(main())
