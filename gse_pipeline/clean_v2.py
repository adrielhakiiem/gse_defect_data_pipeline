"""Clean, audit, and export the CMMS EDR source for review."""

import argparse
import datetime as dt
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import config
import discover_typos_v2
import read_source_v2


_TEXT_FIELDS = (
    "Work Status", "Equipment Code", "Asset Number", "Asset Description",
    "Defects Description", "Work Order Type", "Work Order Trade", "Action Taken",
    "By Month",
)
_NON_DEFECT_PATTERNS = re.compile(
    r"DUTY\s+REPORT|MONITOR\s+STAFF|REQUEST\s+REFUEL(?:L|)ING|GSE\s+STORE",
    re.IGNORECASE,
)
_WORD_RE = re.compile(r"\b[A-Za-z]+\b")


def _stats(changed=0, **extra):
    return {"rows_changed": changed, **extra}


def _text(value):
    return value.strip() if isinstance(value, str) else value


def flag_non_defect_rows(rows):
    """Flag administrative descriptions without removing them."""
    by_type = Counter()
    by_trade = Counter()
    flagged = 0
    for row in rows:
        is_non_defect = bool(_NON_DEFECT_PATTERNS.search(row.get("Defects Description") or ""))
        row["_non_defect_flag"] = is_non_defect
        if is_non_defect:
            flagged += 1
            by_type[row.get("Work Order Type") or "<blank>"] += 1
            by_trade[row.get("Work Order Trade") or "<blank>"] += 1
    print(f"Non-defect rows flagged: {flagged}")
    print(f"  By Work Order Type: {dict(sorted(by_type.items()))}")
    print(f"  By Work Order Trade: {dict(sorted(by_trade.items()))}")
    return _stats(flagged, by_work_order_type=dict(by_type), by_work_order_trade=dict(by_trade))


def normalize_whitespace(rows):
    """Collapse whitespace in descriptions/actions and trim all text fields."""
    changed = 0
    for row in rows:
        before = {field: row.get(field) for field in _TEXT_FIELDS}
        for field in _TEXT_FIELDS:
            value = row.get(field)
            if isinstance(value, str):
                row[field] = re.sub(r"\s+", " ", value).strip() or None
        if before != {field: row.get(field) for field in _TEXT_FIELDS}:
            changed += 1
    print(f"Rows changed by whitespace normalization: {changed}")
    return _stats(changed)


def normalize_categories(rows):
    """Normalize configured display names for categorical fields."""
    mappings = {
        "Work Status": config.WORK_STATUS_MAP,
        "Work Order Trade": config.WORK_ORDER_TRADE_MAP,
    }
    changed = 0
    field_changes = Counter()
    lookup_maps = {
        field: {source.casefold(): target for source, target in mapping.items()}
        for field, mapping in mappings.items()
    }

    for row in rows:
        row_changed = False
        for field, lookup in lookup_maps.items():
            value = row.get(field)
            if not isinstance(value, str):
                continue
            replacement = lookup.get(value.casefold())
            if replacement is not None and replacement != value:
                row[field] = replacement
                field_changes[field] += 1
                row_changed = True
        if row_changed:
            changed += 1

    print(f"Rows changed by category normalization: {changed}")
    print(f"  Work Status values changed: {field_changes['Work Status']}")
    print(f"  Work Order Trade values changed: {field_changes['Work Order Trade']}")
    return _stats(
        changed,
        work_status_values_changed=field_changes["Work Status"],
        work_order_trade_values_changed=field_changes["Work Order Trade"],
    )


def normalize_equipment_code(rows):
    """Normalize both identifiers, removing Excel's trailing ``.0`` values."""
    changed = 0
    field_changes = Counter()

    def normalize(value):
        if isinstance(value, float) and value.is_integer():
            value = str(int(value))
        if value is None:
            return None
        text = str(value).strip()
        whole_number = re.fullmatch(r"([+-]?\d+)\.0+", text)
        if whole_number:
            text = whole_number.group(1)
        return text.upper() or None

    for row in rows:
        before = (row.get("Equipment Code"), row.get("Asset Number"))
        for field in ("Equipment Code", "Asset Number"):
            value = normalize(row.get(field))
            if value != row.get(field):
                field_changes[field] += 1
            row[field] = value
        if before != (row.get("Equipment Code"), row.get("Asset Number")):
            changed += 1

    decimal_rows = []
    for row in rows:
        values = (row.get("Equipment Code"), row.get("Asset Number"))
        if any(isinstance(value, str) and "." in value for value in values):
            decimal_rows.append(row)
    print(f"Rows changed by equipment-code normalization: {changed}")
    print(f"  Equipment Code rows changed: {field_changes['Equipment Code']}")
    print(f"  Asset Number rows changed: {field_changes['Asset Number']}")
    print(f"Rows with a decimal point remaining in an identifier: {len(decimal_rows)}")
    if decimal_rows:
        for row in decimal_rows[:5]:
            print(f"  {row.get('Permit No')}: Equipment Code={row.get('Equipment Code')!r}, "
                  f"Asset Number={row.get('Asset Number')!r}")
    return _stats(
        changed,
        equipment_code_rows_changed=field_changes["Equipment Code"],
        asset_number_rows_changed=field_changes["Asset Number"],
        decimal_identifier_rows=len(decimal_rows),
        decimal_identifier_samples=[
            (row.get("Permit No"), row.get("Equipment Code"), row.get("Asset Number"))
            for row in decimal_rows[:5]
        ],
    )


def _minute_key(row):
    timestamp = row.get("_received_datetime")
    if isinstance(timestamp, dt.datetime):
        return timestamp.replace(second=0, microsecond=0)
    return row.get("Date")


def flag_near_duplicates(rows):
    """Flag repeated asset/time/description-prefix groups for review."""
    groups = defaultdict(list)
    for row in rows:
        description_key = re.sub(r"[^A-Z0-9]", "", (row.get("Defects Description") or "").upper())[:40]
        asset_number = row.get("Asset Number") or row.get("Equipment Code")
        received_minute = _minute_key(row)
        if not asset_number or not received_minute or not description_key:
            continue
        key = (asset_number, received_minute, description_key)
        groups[key].append(row)

    group_count = 0
    rows_affected = 0
    for members in groups.values():
        if len(members) <= 1:
            continue
        group_count += 1
        rows_affected += len(members)
        group_id = f"NEAR-{group_count:04d}"
        for row in members:
            row["_near_dup_group_id"] = group_id
    print(f"Near-duplicate groups: {group_count}; rows affected: {rows_affected}")
    return _stats(rows_affected, groups=group_count)


def _preserve_case(match, replacement):
    original = match.group(0)
    if original.isupper():
        return replacement.upper()
    if original.islower():
        return replacement.lower()
    if original[:1].isupper() and original[1:].islower():
        return replacement[:1].upper() + replacement[1:].lower()
    return replacement


def standardize_spelling(rows):
    """Apply only the explicit, whole-word spelling map to descriptions."""
    substitutions = [
        (re.compile(r"\b" + re.escape(bad) + r"\b", re.IGNORECASE), good)
        for bad, good in config.SPELLING_MAP.items()
    ]
    changed = 0
    for row in rows:
        original = row.get("Defects Description")
        if not isinstance(original, str):
            continue
        cleaned = original
        for pattern, replacement in substitutions:
            cleaned = pattern.sub(lambda match: _preserve_case(match, replacement), cleaned)
        if cleaned != original:
            row["Defects Description"] = cleaned
            changed += 1
    print(f"Rows changed by spelling standardization: {changed}")
    return _stats(changed)


def flag_malay(rows):
    """Flag descriptions containing configured Malay terms."""
    patterns = [re.compile(r"\b" + re.escape(word) + r"\b", re.IGNORECASE)
                for word in config.MALAY_WORDS]
    flagged = 0
    for row in rows:
        description = row.get("Defects Description") or ""
        is_malay = any(pattern.search(description) for pattern in patterns)
        row["_malay_flag"] = is_malay
        flagged += is_malay
    print(f"Rows flagged for Malay terms: {flagged}")
    return _stats(flagged)


def remove_exact_duplicates(rows):
    """Remove exact duplicates after cleaning and return rows plus audit stats."""
    seen = set()
    unique_rows = []
    removed = []
    for row in rows:
        key = tuple((field, repr(row.get(field))) for field in row if not field.startswith("_"))
        if key in seen:
            removed.append(row)
        else:
            seen.add(key)
            unique_rows.append(row)
    print(f"Exact duplicate rows removed: {len(removed)}")
    return unique_rows, _stats(len(removed), removed_rows=removed)


def _display_fields(row):
    return {key: value for key, value in row.items() if key != "_received_datetime"}


def write_review_workbook(rows, stats, output_path):
    """Write cleaned data and auditable review tabs to a local workbook."""
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Cleaned Data"
    fields = list(_display_fields(rows[0]).keys()) if rows else []
    for field in ("_non_defect_flag", "_malay_flag", "_near_dup_group_id"):
        if field not in fields:
            fields.append(field)
    data_sheet.append(fields)
    for cell in data_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5496")
    for row in rows:
        data_sheet.append([_display_fields(row).get(field) for field in fields])

    def add_review_sheet(title, selected_rows):
        sheet = workbook.create_sheet(title)
        sheet.append(fields)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="7F6000")
        for row in selected_rows:
            sheet.append([_display_fields(row).get(field) for field in fields])

    add_review_sheet("Non-Defect Review", [row for row in rows if row.get("_non_defect_flag")])
    add_review_sheet("Near-Duplicate Review", [row for row in rows if row.get("_near_dup_group_id")])
    add_review_sheet("Malay Review", [row for row in rows if row.get("_malay_flag")])
    add_review_sheet("Exact Duplicate Review", stats["exact_duplicates"].get("removed_rows", []))

    summary = workbook.create_sheet("QA Summary", 0)
    summary.append(["Metric", "Value"])

    def append_summary(name, value):
        if isinstance(value, dict):
            for detail, nested_value in value.items():
                if detail == "removed_rows":
                    continue
                append_summary(f"{name}: {detail}", nested_value)
        elif isinstance(value, list):
            for index, item in enumerate(value, start=1):
                append_summary(f"{name} {index}", repr(item))
        else:
            summary.append([name, value])

    for name, value in stats.items():
        append_summary(name, value)
    summary[1][0].font = Font(bold=True)
    summary[1][1].font = Font(bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    workbook.save(output_path)
    print(f"Saved local review workbook: {output_path}")


def clean_rows(rows):
    """Run the independent cleaning steps in an auditable order."""
    stats = {"rows_read": len(rows)}
    stats["whitespace"] = normalize_whitespace(rows)
    stats["categories"] = normalize_categories(rows)
    stats["equipment_code"] = normalize_equipment_code(rows)
    stats["spelling"] = standardize_spelling(rows)
    stats["non_defect"] = flag_non_defect_rows(rows)
    stats["malay"] = flag_malay(rows)
    stats["near_duplicates"] = flag_near_duplicates(rows)
    rows, duplicate_stats = remove_exact_duplicates(rows)
    stats["exact_duplicates"] = duplicate_stats
    stats["rows_exported"] = len(rows)
    return rows, stats


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=config.SOURCE_FILE)
    parser.add_argument("--output", type=Path, default=config.OUTPUT_FILE)
    parser.add_argument("--typo-threshold", type=int, default=8)
    args = parser.parse_args()

    print(f"Loading {args.source}...")
    rows, failed_examples = read_source_v2.read_and_normalize(args.source)
    read_source_v2.print_summary(rows, failed_examples)
    cleaned_rows, stats = clean_rows(rows)
    stats["equipment_code_extraction_failed"] = len(failed_examples)
    stats["missing_or_unparsed_dates"] = sum(1 for row in rows if row.get("Date") is None)
    stats["blank_defect_descriptions"] = sum(
        1 for row in rows if not row.get("Defects Description")
    )
    counter = discover_typos_v2.word_frequencies(
        [row.get("Defects Description") for row in cleaned_rows]
    )
    candidates = discover_typos_v2.find_candidates(counter, args.typo_threshold)
    print(f"Rare-word candidates after cleaning: {len(candidates)}")
    print("Top 40 candidates:")
    for word, count in candidates[:40]:
        print(f"  {count:3d}  {word}")
    write_review_workbook(cleaned_rows, stats, args.output)


if __name__ == "__main__":
    main()
